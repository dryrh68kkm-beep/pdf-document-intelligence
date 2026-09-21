"""Reverse bridge to LP-Tools/expiry-dashboard: given a barcode this pipeline
already OCR'd off a packing list, check whether that barcode is a real,
currently-stocked product according to expiry-dashboard's own daily POS
export (data.csv or data.xlsx, written by that app's fetch-data.ps1) - the
same file its www/index.html auto-loads on startup.

Every live read is merged (upsert only) into a persistent snapshot - see
catalog/expiry_dashboard_snapshot.py's module docstring for why a live
read alone isn't enough: expiry-dashboard's file only ever lists today's
near-expiry items, so a barcode confirmed yesterday can vanish from
today's file without ever having stopped being a real product. A
missing/unreadable live file (the other app not installed, or not run
yet today) falls back to whatever the snapshot has already learned,
rather than losing every match learned so far - this lookup is an extra
confirmation signal, never a hard dependency, but a durable one.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import load_workbook

from pdf_document_intelligence.catalog.expiry_dashboard_snapshot import load_snapshot, merge_into_snapshot

_cache: dict[str, dict[str, str]] = {}
_cache_key: tuple[str, str, float, int] | None = None


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def _read_xlsx_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb[wb.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows_iter, [])]
        return [dict(zip(header, row)) for row in rows_iter]
    finally:
        wb.close()


def _build_index(path: Path) -> dict[str, dict[str, str]]:
    rows = _read_xlsx_rows(path) if path.suffix.lower() in (".xlsx", ".xls") else _read_csv_rows(path)
    index: dict[str, dict[str, str]] = {}
    for row in rows:
        barcode = str(row.get("BAR_CODE") or "").strip()
        if not barcode:
            continue
        index[barcode] = {
            "description": str(row.get("DESCRIPTION") or "").strip(),
            "sub_dept_name": str(row.get("SUB_DEPT_NAME") or "").strip(),
        }
    return index


def load_expiry_dashboard_index(
    path: Path | None, snapshot_path: Path | None = None
) -> dict[str, dict[str, str]]:
    """Cached by (snapshot path, live path, mtime, size) - re-parsed only
    when the file the other app just rewrote actually differs from what's
    cached, so a file that hasn't changed today costs nothing beyond one
    stat() call. When there's no live file to read (missing today, or the
    other app was never pointed at one), returns the persistent snapshot
    as-is - every barcode ever confirmed, never wiped just because today's
    read didn't happen."""
    global _cache, _cache_key
    if path is not None:
        try:
            stat = path.stat()
        except OSError:
            path = None
        else:
            key = (str(snapshot_path), str(path), stat.st_mtime, stat.st_size)
            if key == _cache_key:
                return _cache
            try:
                live_entries = _build_index(path)
            except (OSError, csv.Error, KeyError, ValueError):
                path = None
            else:
                _cache = merge_into_snapshot(live_entries, snapshot_path)
                _cache_key = key
                return _cache
    # No usable live file this call - serve whatever the persistent
    # snapshot already has rather than nothing.
    return load_snapshot(snapshot_path)


def resolve_data_path(www_dir: str) -> Path | None:
    """expiry-dashboard's own tryAutoLoad() tries these two names in this
    same order (fetch-data.ps1 writes whichever matches the store's source
    file type) - mirror that instead of assuming one or the other."""
    if not www_dir:
        return None
    base = Path(www_dir)
    for name in ("data.csv", "data.xlsx"):
        candidate = base / name
        if candidate.is_file():
            return candidate
    return None


def lookup_barcode(
    barcode: str | None, path: Path | None, snapshot_path: Path | None = None
) -> dict[str, str] | None:
    if not barcode:
        return None
    return load_expiry_dashboard_index(path, snapshot_path).get(str(barcode))
