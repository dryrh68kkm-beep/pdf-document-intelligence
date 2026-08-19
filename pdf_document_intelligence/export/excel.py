"""ExportService: 4-sheet Excel export (Data / Validation Issues / Document
Summary / Audit). Code fields are written with an explicit text number
format so leading zeros survive (proposal §23/§56); numeric fields keep
their native type so Excel can sum them.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from pdf_document_intelligence.models.document import DocumentResult
from pdf_document_intelligence.templates.packing_list_bigc import COLUMNS

_HEADER_FONT = Font(bold=True)


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT


def _write_data_sheet(ws: Worksheet, result: DocumentResult) -> None:
    headers = ["department", *[c.canonical_name for c in COLUMNS], "row_confidence", "review_required"]
    _write_header(ws, headers)
    for table in result.tables:
        for row in table.rows:
            values: list[object] = [table.name]
            for col in COLUMNS:
                fv = row.fields.get(col.canonical_name)
                values.append(fv.value if fv else None)
            review = any(f.review_required for f in row.fields.values())
            values.append(row.confidence_band)
            values.append("YES" if review else "")
            ws.append(values)
    # Code columns (dn_no, do_no, order_no, pallet, lot, article, barcode)
    # are written as text so Excel never strips leading zeros.
    code_col_indices = [i + 2 for i, c in enumerate(COLUMNS) if c.field_type == "code"]  # +2: 1-indexed, +department col
    for col_idx in code_col_indices:
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=col_idx).number_format = "@"


def _write_issues_sheet(ws: Worksheet, result: DocumentResult) -> None:
    _write_header(ws, ["severity", "code", "table", "row_index", "field", "message"])
    for issue in [*result.validation.errors, *result.validation.warnings]:
        ws.append([issue.severity, issue.code, issue.table, issue.row_index, issue.field, issue.message])


def _write_summary_sheet(ws: Worksheet, result: DocumentResult) -> None:
    _write_header(ws, ["key", "value"])
    total_rows = sum(len(t.rows) for t in result.tables)
    low_conf_rows = sum(1 for t in result.tables for r in t.rows if r.confidence_band == "LOW")
    rows_needing_review = sum(
        1 for t in result.tables for r in t.rows if any(f.review_required for f in r.fields.values())
    )
    summary = [
        ("filename", result.filename),
        ("pages", result.pages),
        ("document_type", result.document_type),
        ("tables", len(result.tables)),
        ("total_rows", total_rows),
        ("low_confidence_rows", low_conf_rows),
        ("rows_needing_review", rows_needing_review),
        ("validation_errors", len(result.validation.errors)),
        ("validation_warnings", len(result.validation.warnings)),
        ("reconciliation_status", "PASSED" if result.validation.reconciled else "FAILED"),
        ("document_confidence_pct", round(result.confidence, 2)),
        ("status", result.status),
        ("engine_version", result.engine_version),
        ("parser_version", result.parser_version),
        ("template_version", result.template_version),
    ]
    for key, value in summary:
        ws.append([key, value])


def _write_audit_sheet(ws: Worksheet, result: DocumentResult) -> None:
    _write_header(ws, ["step", "detail", "timestamp", "duration_ms"])
    for entry in result.processing_log:
        ws.append([entry.step, entry.detail, entry.timestamp.isoformat(), entry.duration_ms])


def export_to_excel(result: DocumentResult, output_path: Path) -> Path:
    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "Extracted Data"
    _write_data_sheet(data_ws, result)

    issues_ws = wb.create_sheet("Validation Issues")
    _write_issues_sheet(issues_ws, result)

    summary_ws = wb.create_sheet("Document Summary")
    _write_summary_sheet(summary_ws, result)

    audit_ws = wb.create_sheet("Audit")
    _write_audit_sheet(audit_ws, result)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path
