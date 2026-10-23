"""Expiry feed: the same data.csv / status.json contract as the standalone
expiry-dashboard (fetch-data.ps1), so the dashboard page, the PDF pipeline's
reverse lookup and the sync hub all read one file."""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import threading
from datetime import datetime, timezone
from pathlib import Path

from store_suite.paths import expiry_dir

REQUIRED_COLUMNS = ("BAR_CODE", "DESCRIPTION", "STOCK_QTY", "DAY_LEFT", "SUB_DEPT_NAME", "DIV", "DEPT")
_write_lock = threading.Lock()
_cache: dict = {"key": None, "rows": []}


def data_file() -> Path | None:
    for name in ("data.csv", "data.xlsx"):
        p = expiry_dir() / name
        if p.is_file():
            return p
    return None


def _read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
            header = [str(h or "").strip() for h in next(it, [])]
            return [{h: ("" if v is None else v) for h, v in zip(header, row)} for row in it]
        finally:
            wb.close()
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def validate(path: Path) -> tuple[bool, str, int]:
    """(ok, reason, rowCount) - same checks as fetch-data.ps1's Test-CsvValid."""
    try:
        if path.stat().st_size <= 0:
            return False, "ไฟล์มีขนาด 0 ไบต์", 0
        rows = _read_rows(path)
    except Exception as exc:  # noqa: BLE001 - any unreadable file is just "invalid"
        return False, f"อ่านไฟล์ไม่ได้: {exc}", 0
    if not rows:
        return False, "ไฟล์มีแค่ header ไม่มีข้อมูล", 0
    missing = [c for c in REQUIRED_COLUMNS if c not in rows[0]]
    if missing:
        return False, "ไฟล์ขาดคอลัมน์ที่จำเป็น: " + ", ".join(missing), 0
    return True, "", len(rows)


def load_rows() -> list[dict]:
    """Parsed rows of the current data file, cached until it changes."""
    path = data_file()
    if path is None:
        return []
    st = path.stat()
    key = (str(path), st.st_mtime_ns, st.st_size)
    if _cache["key"] != key:
        try:
            _cache["rows"] = _read_rows(path)
        except Exception:  # noqa: BLE001
            _cache["rows"] = []
        _cache["key"] = key
    return _cache["rows"]


def read_status() -> dict:
    try:
        return json.loads((expiry_dir() / "status.json").read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return {}


def _write_status(success: bool, name: str, size: int, rows: int, error: str, src_mtime: str | None, src_len: int | None) -> None:
    status = {
        "lastAttempt": datetime.now(timezone.utc).astimezone().isoformat(),
        "success": success,
        "fileName": name,
        "fileSizeBytes": size,
        "rowCount": rows,
        "errorMessage": error,
        "sourceLastWriteTimeUtc": src_mtime,
        "sourceLength": src_len,
    }
    tmp = expiry_dir() / "status.json.tmp"
    tmp.write_text(json.dumps(status, ensure_ascii=False), encoding="utf-8")
    tmp.replace(expiry_dir() / "status.json")


def install_file(src: Path, display_name: str | None = None) -> dict:
    """Validate `src` and, only if good, atomically replace data.csv/xlsx.
    A bad file never touches the existing data (status.json says why)."""
    name = display_name or src.name
    with _write_lock:
        ok, reason, rows = validate(src)
        st = src.stat()
        src_mtime = datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat()
        if not ok:
            _write_status(False, name, st.st_size, 0, reason, src_mtime, st.st_size)
            return {"ok": False, "error": reason}
        ext = ".xlsx" if src.suffix.lower() in (".xlsx", ".xlsm") else ".csv"
        dst = expiry_dir() / f"data{ext}"
        tmp = expiry_dir() / f"data{ext}.tmp"
        shutil.copyfile(src, tmp)
        os.replace(tmp, dst)
        other = expiry_dir() / ("data.xlsx" if ext == ".csv" else "data.csv")
        other.unlink(missing_ok=True)  # the page tries csv first - never leave a stale twin
        _write_status(True, name, st.st_size, rows, "", src_mtime, st.st_size)
    return {"ok": True, "rows": rows, "file": name}


def fetch_from_source(cfg: dict) -> dict:
    """Pull the newest NEARLY_EXPIRED file from the configured network folder.
    Skips work when the source is unchanged since the last good copy."""
    folder = (cfg.get("expirySourceFolder") or "").strip()
    if not folder:
        return {"ok": True, "skipped": "no source folder configured"}
    root = Path(folder)
    try:
        if not root.is_dir():
            raise OSError("ไม่พบโฟลเดอร์ต้นทาง")
        base = cfg.get("expirySourceBasename") or "NEARLY_EXPIRED"
        exts = [e.lower() for e in cfg.get("expirySourceExtensions") or [".csv"]]
        candidates = [p for e in exts for p in [root / f"{base}{e}", root / f"{base}{e.upper()}"] if p.is_file()]
        if not candidates:
            prefix = base.rsplit("_", 1)[0].lower() if base.count("_") >= 2 else base.lower()
            candidates = [p for p in root.iterdir() if p.is_file() and p.name.lower().startswith(prefix) and p.suffix.lower() in exts]
        if not candidates:
            raise OSError(f"ไม่พบไฟล์ {base}* ในโฟลเดอร์ต้นทาง")
        newest = max(candidates, key=lambda p: p.stat().st_mtime)
    except OSError as exc:
        with _write_lock:
            prev = read_status()
            _write_status(False, prev.get("fileName", ""), prev.get("fileSizeBytes", 0), prev.get("rowCount", 0), str(exc), prev.get("sourceLastWriteTimeUtc"), prev.get("sourceLength"))
        return {"ok": False, "error": str(exc)}
    st = newest.stat()
    prev = read_status()
    src_mtime = datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat()
    if prev.get("success") and prev.get("sourceLastWriteTimeUtc") == src_mtime and prev.get("sourceLength") == st.st_size and data_file():
        return {"ok": True, "skipped": "source unchanged"}
    return install_file(newest)
